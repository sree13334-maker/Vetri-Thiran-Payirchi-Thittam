from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from fpdf import FPDF

from utils.text_utils import sanitize_text


BASE_DIR = Path(
    __file__
).resolve().parent.parent


LOGO_PATH = (
    BASE_DIR /
    "assets" /
    "logo.png"
)


def format_txt(text):

    clean_text = sanitize_text(
        text
    )

    return clean_text.encode(
        "utf-8"
    )


def add_logo_to_docx(document):

    if LOGO_PATH.exists():

        paragraph = (
            document.add_paragraph()
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            str(LOGO_PATH),
            width=Inches(1.25)
        )


def format_docx(
    text,
    doc_type="Legal Document"
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)

    section.bottom_margin = Inches(0.7)

    section.left_margin = Inches(0.8)

    section.right_margin = Inches(0.8)


    add_logo_to_docx(
        document
    )


    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        sanitize_text(
            doc_type
        ).upper()
    )

    title_run.bold = True

    title_run.font.name = (
        "Times New Roman"
    )

    title_run.font.size = Pt(16)


    clean_text = sanitize_text(
        text
    )


    for raw_line in clean_text.splitlines():

        line = raw_line.strip()

        if not line:
            continue


        paragraph = (
            document.add_paragraph()
        )

        paragraph.paragraph_format.space_after = (
            Pt(7)
        )


        is_heading = (

            len(line) <= 90

            and (

                line.isupper()

                or line.endswith(":")

                or line.lower()
                in {
                    "parties",
                    "effective date",
                    "definitions",
                    "signatures",
                    "important notice"
                }
            )
        )


        run = paragraph.add_run(
            line
        )

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(11)


        if is_heading:

            run.bold = True


    footer = (
        section
        .footer
        .paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review before use."
    )

    footer_run.font.name = (
        "Times New Roman"
    )

    footer_run.font.size = Pt(8)


    output = BytesIO()

    document.save(
        output
    )

    return output.getvalue()


class LegalEasePDF(FPDF):

    def __init__(
        self,
        doc_type
    ):

        super().__init__()

        self.doc_type = doc_type


    def header(self):

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=90,
                    y=8,
                    w=30
                )

                self.ln(22)

            except Exception:

                self.ln(8)

        else:

            self.ln(8)


        self.set_font(
            "Helvetica",
            "B",
            14
        )

        self.cell(
            0,
            8,
            self.doc_type[:80],
            new_x="LMARGIN",
            new_y="NEXT",
            align="C"
        )

        self.ln(4)


    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=8
        )

        self.cell(
            0,
            8,
            (
                "LegalEase - AI-generated draft | "
                f"Page {self.page_no()}"
            ),
            align="C"
        )


def format_pdf(
    text,
    doc_type="Legal Document"
):

    pdf = LegalEasePDF(
        doc_type
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()


    pdf.set_font(
        "Helvetica",
        size=11
    )


    clean_text = sanitize_text(
        text
    )


    for raw_line in clean_text.splitlines():

        line = raw_line.strip()


        if not line:

            pdf.ln(3)

            continue


        is_heading = (

            len(line) <= 90

            and (

                line.isupper()

                or line.endswith(":")

                or line.lower()
                in {
                    "parties",
                    "effective date",
                    "definitions",
                    "signatures",
                    "important notice"
                }
            )
        )


        if is_heading:

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                7,
                line
            )

            pdf.set_font(
                "Helvetica",
                size=11
            )

        else:

            pdf.multi_cell(
                0,
                6,
                line
            )


        pdf.ln(1)


    return bytes(
        pdf.output()
    )